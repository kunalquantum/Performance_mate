"""
Detect what kind of workbook was uploaded.

Two known types today:
    "weekly_esp"  - the weekly ESP campaign report
                    (Summary / Opened 3+ Times / Clicked 1+ Times)
    "crm_sync"    - a Zoho CRM sync dump
                    (Accounts / Leads / Contacts / masterlist)

Returns "unknown" if neither matches. Detection is signature-based:
sheet names + header keywords, so slight column drift is tolerated.
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import openpyxl


def _norm(s) -> str:
    return str(s or "").strip().lower()


def _sheet_headers(ws) -> set:
    """Return the first row of a sheet as a lower-cased set of tokens."""
    row = next(ws.iter_rows(values_only=True), None) or ()
    return {_norm(c) for c in row if c is not None}


def detect(path: str | Path) -> Dict[str, str | dict]:
    """
    Return a dict:
        {"type": "weekly_esp" | "crm_sync" | "unknown",
         "reason": short human explanation,
         "signals": {<what matched>}}
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets_lc = [_norm(n) for n in wb.sheetnames]
    signals = {"sheets": sheets_lc}

    # ---- Weekly ESP signature ----------------------------------------
    esp_signals = 0
    if any("summary" == s for s in sheets_lc):
        esp_signals += 1
    if any("opened" in s for s in sheets_lc):
        esp_signals += 1
    if any("clicked" in s for s in sheets_lc):
        esp_signals += 1
    if esp_signals >= 2:
        # Confirm with header content on the Summary sheet
        summary_name = next((n for n in wb.sheetnames
                              if _norm(n) == "summary"), None)
        if summary_name:
            ws = wb[summary_name]
            # First 10 rows might contain the header
            saw_campaign_header = False
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i > 10:
                    break
                toks = {_norm(c) for c in row if c is not None}
                if "campaign name" in toks and (
                        "opened" in toks or "opened (%)" in toks
                        or any("opened" in t for t in toks)):
                    saw_campaign_header = True
                    break
            if saw_campaign_header:
                return {"type": "weekly_esp",
                         "reason": ("Sheet names + Summary header match "
                                     "the weekly ESP shape."),
                         "signals": signals}

    # ---- CRM sync signature ------------------------------------------
    crm_sheets = {"accounts", "leads", "contacts", "masterlist"}
    matched = crm_sheets & set(sheets_lc)
    if len(matched) >= 2:
        # Confirm with a header check on one of them
        target = None
        for candidate in ("contacts", "leads", "accounts"):
            if candidate in sheets_lc:
                target = wb.sheetnames[sheets_lc.index(candidate)]
                break
        if target:
            hdrs = _sheet_headers(wb[target])
            zoho_hallmarks = {"email", "first_name", "last_name",
                               "account_name", "modified_time",
                               "created_time", "lead_source",
                               "lead_status", "phone"}
            if len(hdrs & zoho_hallmarks) >= 3:
                return {"type": "crm_sync",
                         "reason": (f"Sheets {sorted(matched)} + Zoho "
                                     f"-style headers detected."),
                         "signals": {**signals, "matched_sheets":
                                       sorted(matched)}}

    return {"type": "unknown",
             "reason": ("Neither weekly ESP nor CRM sync signature "
                          "matched."),
             "signals": signals}


if __name__ == "__main__":
    import sys
    p = sys.argv[1] if len(sys.argv) > 1 else \
        r"C:\Users\KunalSharma\Downloads\30Aug-4Sept.xlsx"
    r = detect(p)
    print(r)
